"""Persist experiment histories and prepare the submission table.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    output_dir = Path(results_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    exp_id = str(result["cfg"]["exp_id"])
    payload = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }

    def json_default(value):
        if hasattr(value, "tolist"):
            return value.tolist()
        if hasattr(value, "item"):
            return value.item()
        if isinstance(value, Path):
            return str(value)
        raise TypeError(f"Object of type {type(value).__name__} is not JSON serializable")

    path = output_dir / f"{exp_id}.json"
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=json_default), encoding="utf-8")
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.exists():
        return []
    results = [json.loads(path.read_text(encoding="utf-8")) for path in sorted(directory.glob("*.json"))]
    return sorted(results, key=lambda item: item["cfg"].get("exp_id", ""))


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result["cfg"]
    summary = result["summary"]
    row = {**cfg, **summary}
    row["hidden"] = "-".join(str(width) for width in cfg.get("hidden", ()))
    row["eval_acc"] = None if eval_scores is None else eval_scores.get("accuracy")
    row["eval_macro_f1"] = None if eval_scores is None else eval_scores.get("macro_f1")
    row["figure_file"] = f"figures/{cfg['exp_id']}.png"
    row["notes"] = notes or cfg.get("notes", "")
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_exp_ids: list[str] | None = None,
               summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(template_path)
    worksheet = workbook["Experiments"]
    headers = [cell.value for cell in worksheet[1]]
    formula_columns = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            if header in formula_columns:
                continue
            if header in row:
                value = row[header]
                if isinstance(value, (dict, list, tuple)):
                    value = json.dumps(value, ensure_ascii=False)
                worksheet.cell(row=row_index, column=column_index, value=value)

    # The template's Seeds and Summary sheets contain formulas and their own
    # headers. Fill only the input cells so those formulas remain intact.
    if seed_exp_ids is not None and "Seeds" in workbook.sheetnames:
        seeds = workbook["Seeds"]
        seed_headers = [cell.value for cell in seeds[1]]
        exp_id_column = next(
            index for index, header in enumerate(seed_headers, start=1)
            if isinstance(header, str) and header.startswith("exp_id")
        )
        for offset, exp_id in enumerate(seed_exp_ids, start=2):
            seeds.cell(row=offset, column=exp_id_column, value=exp_id)

    if summary_notes and "Summary" in workbook.sheetnames:
        summary = workbook["Summary"]
        summary_headers = [cell.value for cell in summary[1]]
        group_column = summary_headers.index("group") + 1
        notes_column = next(
            index for index, header in enumerate(summary_headers, start=1)
            if isinstance(header, str) and header.startswith("nhận xét")
        )
        for row_index in range(2, summary.max_row + 1):
            group = summary.cell(row=row_index, column=group_column).value
            if group in summary_notes:
                summary.cell(row=row_index, column=notes_column, value=summary_notes[group])
    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
